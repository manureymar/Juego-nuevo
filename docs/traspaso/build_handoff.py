"""Render the adjacent handoff Markdown as a readable DOCX using python-docx.

Run from any directory with Python and python-docx installed.
The Markdown is the editable content source. Generated QA pages are not published.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / 'Robot-Pulse-Traspaso.md'
OUT = ROOT / 'Robot-Pulse-Traspaso.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin, sec.bottom_margin = Inches(.68), Inches(.68)
sec.left_margin, sec.right_margin = Inches(.72), Inches(.72)
sec.header_distance, sec.footer_distance = Inches(.28), Inches(.28)
usable = 8.27 - .72 * 2

for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet', 'List Number', 'Header', 'Footer']:
    s = doc.styles[name]
    s.font.name = 'DejaVu Sans'
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.font.size = Pt(10)
    s.paragraph_format.space_after = Pt(6)
    s.paragraph_format.line_spacing = 1.08
    s.paragraph_format.widow_control = True
    for el in s.element.xpath('.//w:color'):
        for attr in ['themeColor','themeTint','themeShade']:
            el.attrib.pop(qn('w:'+attr), None)
doc.styles['Title'].font.size = Pt(24)
doc.styles['Title'].font.bold = True
doc.styles['Title'].paragraph_format.space_after = Pt(12)
for name, size in [('Heading 1',16), ('Heading 2',12), ('Heading 3',11)]:
    st = doc.styles[name]
    st.font.size = Pt(size)
    st.font.bold = True
    st.paragraph_format.space_before = Pt(13)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.keep_with_next = True
    st.paragraph_format.keep_together = True
for name in ['List Bullet','List Number']:
    doc.styles[name].paragraph_format.space_after = Pt(4)
for el in doc.styles.element.xpath('.//w:pBdr'):
    el.getparent().remove(el)
doc.styles['Footer'].font.size = Pt(8)
doc.core_properties.title = 'Traspaso completo de Robot Pulse'
doc.core_properties.subject = 'Continuidad de producto, diseño y desarrollo Android 0.6.0'
doc.core_properties.author = 'Robot Pulse'
doc.core_properties.keywords = 'Robot Pulse, traspaso, minería, Android, GitHub'

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Robot Pulse 0.6.0 · Traspaso · ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

def link(p, text, url):
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    r = OxmlElement('w:r'); rp = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '165A84'); rp.append(color)
    size = OxmlElement('w:sz'); size.set(qn('w:val'), '18'); rp.append(size)
    r.append(rp); t = OxmlElement('w:t'); t.text = text; r.append(t); h.append(r); p._p.append(h)

def inline(p, text, table=False):
    pieces = re.split(r'(\*\*.*?\*\*|`[^`]+`|https?://[^\s]+)', text)
    for s in pieces:
        if not s: continue
        if s.startswith('http'):
            link(p, s, s)
            continue
        r = p.add_run(s[2:-2] if s.startswith('**') else s[1:-1] if s.startswith('`') else s)
        if s.startswith('**'): r.bold = True
        if s.startswith('`'):
            r.font.name = 'DejaVu Sans Mono'
            r.font.size = Pt(8.0 if table else 8.7)
        elif table:
            r.font.size = Pt(9)

def add_table(rows):
    rows = [r for r in rows if not all(re.fullmatch(r':?-+:?', x.strip()) for x in r)]
    n = len(rows[0]); table = doc.add_table(rows=0, cols=n)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    # File maps need generous path space; numbers/prices need much less.
    is_map = any(x in rows[0][0].lower() for x in ['ruta','archivo','carpeta'])
    if n == 2: widths = [usable*.49, usable*.51] if is_map else [usable*.34, usable*.66]
    else: widths = [usable*.24, usable*.50, usable*.26]
    for col,w in zip(table.columns,widths): col.width = Inches(w)
    pr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e = OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
    pr.append(borders)
    for idx, row in enumerate(rows):
        cells = table.add_row().cells
        trpr = table.rows[-1]._tr.get_or_add_trPr()
        cant = OxmlElement('w:cantSplit'); trpr.append(cant)
        if idx == 0:
            repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
        for ci, (cell, value) in enumerate(zip(cells,row)):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp = cell._tc.get_or_add_tcPr()
            mar = OxmlElement('w:tcMar')
            for side, val in [('top','80'),('bottom','80'),('left','95'),('right','95')]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),val);e.set(qn('w:type'),'dxa');mar.append(e)
            cp.append(mar)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE7EF' if idx==0 else 'FFFFFF');cp.append(shade)
            p=cell.paragraphs[0]
            p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=1.05
            inline(p,value,True)
            if idx==0:
                p.paragraph_format.keep_with_next=True
                for r in p.runs:r.bold=True;r.font.color.rgb=RGBColor(0,0,0)
            elif not is_map and (value.replace('.','').replace(',','').replace('$','').isdigit()):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    spacer=doc.add_paragraph();spacer.paragraph_format.space_after=Pt(0);spacer.paragraph_format.space_before=Pt(0);spacer.paragraph_format.line_spacing=Pt(3);spacer.add_run(' ').font.size=Pt(3)

lines = SOURCE.read_text().splitlines()
i=0
while i<len(lines):
    line=lines[i]
    if not line.strip():i+=1;continue
    if line.startswith('```'):
        i+=1;code=[]
        while i<len(lines) and not lines[i].startswith('```'):code.append(lines[i]);i+=1
        p=doc.add_paragraph()
        p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(8)
        p.paragraph_format.line_spacing=1.05
        for j,s in enumerate(code):
            r=p.add_run(('\n' if j else '')+s);r.font.name='DejaVu Sans Mono';r.font.size=Pt(8.5)
        i+=1;continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            rows.append([x.strip() for x in lines[i].strip().strip('|').split('|')]);i+=1
        add_table(rows);continue
    if line.startswith('# '):
        p=doc.add_paragraph(line[2:],'Title')
    elif line.startswith('## '):
        p=doc.add_paragraph(line[3:],'Heading 1')
    elif line.startswith('### '):
        p=doc.add_paragraph(line[4:],'Heading 2')
    elif line.startswith('- '):
        p=doc.add_paragraph(style='List Bullet');inline(p,line[2:])
    elif re.match(r'^\d+\. ',line):
        # Keep explicit numbering instead of a global Word sequence.
        p=doc.add_paragraph();p.paragraph_format.left_indent=Inches(.2);p.paragraph_format.first_line_indent=Inches(-.2);inline(p,line)
    elif line.startswith('> '):
        p=doc.add_paragraph();p.paragraph_format.left_indent=Inches(.15);inline(p,line[2:])
    else:
        p=doc.add_paragraph();inline(p,line)
    i+=1

for el in doc.element.xpath('.//w:pBdr'):
    el.getparent().remove(el)
doc.save(OUT)
print(OUT)
